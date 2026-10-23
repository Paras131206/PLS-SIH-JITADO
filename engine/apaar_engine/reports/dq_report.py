"""engine/apaar_engine/reports/dq_report.py

Batch data quality report exporter for APAAR-SDMS:
- BatchReport -> XLSX (Summary, Issues, Intake, Column Metrics sheets)
- BatchReport -> JSON

Owned by Team A (Data & Intelligence).
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Optional, Union

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from apaar_engine.models import BatchReport


def style_header_row(ws, num_cols: int, fill_color: str = "1E3A8A") -> None:
    """Style sheet header with bold white text on navy blue background."""
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def auto_fit_columns(ws) -> None:
    """Adjust worksheet column widths based on contents."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def export_dq_report_xlsx(
    report: BatchReport,
    output_path: Optional[Union[str, Path]] = None,
) -> bytes:
    """Export BatchReport to an Excel XLSX workbook with 4 sheets:

    1. Summary
    2. Issues
    3. Intake
    4. Column Metrics
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Metric", "Value"])
    style_header_row(ws_summary, 2)
    ws_summary.append(["Batch ID", report.batch_id])
    ws_summary.append(["Institution AICTE ID", report.institution_aicte_id])
    ws_summary.append(["Total Rows", report.rows_total])
    ws_summary.append(["Valid Rows", report.rows_valid])
    ws_summary.append(["Warning Rows", report.rows_warning])
    ws_summary.append(["Error Rows", report.rows_error])
    ws_summary.append(["Quality Score", f"{report.quality_score:.2f}%"])
    auto_fit_columns(ws_summary)

    # Sheet 2: Issues
    ws_issues = wb.create_sheet(title="Issues")
    ws_issues.append(["Row", "Field", "Rule ID", "Severity", "Message", "Masked Value"])
    style_header_row(ws_issues, 6)
    for iss in report.issues:
        ws_issues.append(
            [
                iss.row if iss.row is not None else "-",
                iss.field or "-",
                iss.rule_id,
                iss.severity.value,
                iss.message,
                iss.value_masked or "-",
            ]
        )
    auto_fit_columns(ws_issues)

    # Sheet 3: Intake
    ws_intake = wb.create_sheet(title="Intake")
    ws_intake.append(
        [
            "Program Code",
            "Approved Intake",
            "Total Enrolled",
            "Regular Enrolled",
            "Supernumerary Enrolled",
            "Excess Enrolment",
            "Supernumerary Breakdown",
        ]
    )
    style_header_row(ws_intake, 7)
    for intake in report.intake_summary:
        breakdown_str = (
            json.dumps(intake.supernumerary_by_scheme) if intake.supernumerary_by_scheme else "-"
        )
        ws_intake.append(
            [
                intake.program_code,
                intake.approved,
                intake.enrolled,
                intake.enrolled_regular,
                intake.supernumerary,
                intake.excess,
                breakdown_str,
            ]
        )
    auto_fit_columns(ws_intake)

    # Sheet 4: Column Metrics
    ws_metrics = wb.create_sheet(title="Column Metrics")
    ws_metrics.append(["Column Name", "Completeness Rate", "Uniqueness Rate"])
    style_header_row(ws_metrics, 3)
    comp_map = report.metrics.get("completeness", {})
    uniq_map = report.metrics.get("uniqueness", {})
    all_cols = sorted(set(comp_map.keys()) | set(uniq_map.keys()))
    for col_name in all_cols:
        c_rate = comp_map.get(col_name, 0.0)
        u_rate = uniq_map.get(col_name, 0.0)
        ws_metrics.append(
            [
                col_name,
                f"{c_rate * 100.0:.2f}%",
                f"{u_rate * 100.0:.2f}%",
            ]
        )
    auto_fit_columns(ws_metrics)

    # Save output
    stream = io.BytesIO()
    wb.save(stream)
    content = stream.getvalue()

    if output_path is not None:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "wb") as f:
            f.write(content)

    return content


def export_dq_report_json(
    report: BatchReport,
    output_path: Optional[Union[str, Path]] = None,
    indent: int = 2,
) -> str:
    """Export BatchReport to JSON formatted string."""
    json_str = report.model_dump_json(indent=indent)
    if output_path is not None:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(json_str)
    return json_str
