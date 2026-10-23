"""engine/tests/test_dq_report.py

Tests for BatchReport exporter (JSON and XLSX).
Verifies:
- 4 sheets created in XLSX: Summary, Issues, Intake, Column Metrics
- JSON export produces valid, readable BatchReport payload

Owned by Team A (Data & Intelligence).
"""

import io
import json
from pathlib import Path

import openpyxl

from apaar_engine.models import (
    BatchReport,
    IntakeSummary,
    Issue,
    Severity,
)
from apaar_engine.reports.dq_report import export_dq_report_json, export_dq_report_xlsx


def create_sample_batch_report() -> BatchReport:
    return BatchReport(
        batch_id="b7f1418b-5975-4d76-8fe9-2b0244439162",
        institution_aicte_id="1-10001",
        rows_total=100,
        rows_valid=95,
        rows_warning=2,
        rows_error=3,
        quality_score=94.5,
        metrics={
            "completeness": {"apaar_id": 1.0, "dob": 0.98},
            "uniqueness": {"apaar_id": 0.99, "institution_student_id": 1.0},
        },
        issues=[
            Issue(
                row=12,
                field="apaar_id",
                rule_id="V002",
                severity=Severity.ERROR,
                message="apaar_id missing or not 12 digits",
                value_masked="1234****9012",
            ),
            Issue(
                row=15,
                field="dob",
                rule_id="V016",
                severity=Severity.WARN,
                message="Student age 8 implausible for UG",
                value_masked=None,
            ),
        ],
        intake_summary=[
            IntakeSummary(
                program_code="UG-CSE",
                approved=75,
                enrolled=70,
                supernumerary=3,
                excess=0,
                enrolled_regular=67,
                supernumerary_by_scheme={"TFW": 3},
            )
        ],
    )


def test_export_dq_report_json(tmp_path: Path):
    report = create_sample_batch_report()
    json_path = tmp_path / "report.json"
    json_str = export_dq_report_json(report, output_path=json_path)

    assert json_path.exists()
    parsed = json.loads(json_str)
    assert parsed["batch_id"] == "b7f1418b-5975-4d76-8fe9-2b0244439162"
    assert len(parsed["issues"]) == 2
    assert parsed["quality_score"] == 94.5


def test_export_dq_report_xlsx(tmp_path: Path):
    report = create_sample_batch_report()
    xlsx_path = tmp_path / "report.xlsx"
    raw_bytes = export_dq_report_xlsx(report, output_path=xlsx_path)

    assert len(raw_bytes) > 0
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))
    sheet_names = wb.sheetnames
    assert "Summary" in sheet_names
    assert "Issues" in sheet_names
    assert "Intake" in sheet_names
    assert "Column Metrics" in sheet_names

    # Check Summary sheet content
    ws_summary = wb["Summary"]
    assert ws_summary["A2"].value == "Batch ID"
    assert ws_summary["B2"].value == "b7f1418b-5975-4d76-8fe9-2b0244439162"

    # Check Issues sheet content
    ws_issues = wb["Issues"]
    assert ws_issues["A2"].value == 12
    assert ws_issues["C2"].value == "V002"
    assert ws_issues["D2"].value == "ERROR"

    # Check Intake sheet content
    ws_intake = wb["Intake"]
    assert ws_intake["A2"].value == "UG-CSE"
    assert ws_intake["B2"].value == 75
    assert ws_intake["D2"].value == 67

    # Check Column Metrics sheet content
    ws_metrics = wb["Column Metrics"]
    metric_cols = [row[0] for row in ws_metrics.iter_rows(min_row=2, max_col=1, values_only=True)]
    assert "apaar_id" in metric_cols
    assert "dob" in metric_cols
