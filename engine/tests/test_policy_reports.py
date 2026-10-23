"""engine/tests/test_policy_reports.py

Unit tests for policy aggregate reports and small-cell suppression (Phase 4, Team A).
Verifies:
- Small-cell suppression (<5 -> '<5') on counts for privacy preservation.
- Aggregation by state, program, category, quota, and year.
- Multi-sheet Excel workbook export with proper styling and structure.
"""

from pathlib import Path

import openpyxl
import pandas as pd

from apaar_engine.reports.policy_report import (
    apply_small_cell_suppression,
    build_policy_aggregates,
    export_policy_report_xlsx,
)


def test_apply_small_cell_suppression():
    raw_df = pd.DataFrame(
        {
            "Category": ["GEN", "OBC-NCL", "SC", "ST", "EWS"],
            "Count": [15, 4, 1, 0, 8],
            "Supernumerary": [0, 2, 0, 5, 3],
        }
    )

    suppressed = apply_small_cell_suppression(raw_df, threshold=5, suppress_label="<5")

    # Counts 1, 2, 3, 4 must become '<5'
    # Count 0 should remain 0
    # Counts >= 5 should remain unchanged
    assert suppressed["Count"].tolist() == [15, "<5", "<5", 0, 8]
    assert suppressed["Supernumerary"].tolist() == [0, "<5", 0, 5, "<5"]


def test_build_policy_aggregates_synthetic_data():
    repo_root = Path(__file__).resolve().parents[2]
    students_csv = repo_root / "data" / "synthetic" / "students.csv"
    students_df = pd.read_csv(students_csv)

    aggregates = build_policy_aggregates(students_df, suppress=True, threshold=5)

    # Must contain all required policy aggregation dimensions
    expected_tables = {
        "by_state",
        "by_program",
        "by_category",
        "by_quota",
        "by_year",
        "state_x_program",
        "category_x_quota",
    }
    assert set(aggregates.keys()) == expected_tables

    df_state = aggregates["by_state"]
    assert not df_state.empty
    assert "State" in df_state.columns
    assert "Total Students" in df_state.columns
    # Check states are mapped
    states = df_state["State"].tolist()
    assert "Karnataka" in states
    assert "Maharashtra" in states

    df_prog = aggregates["by_program"]
    assert not df_prog.empty
    assert "Program" in df_prog.columns
    programs = df_prog["Program"].tolist()
    assert "UG-CSE" in programs

    df_cat = aggregates["by_category"]
    assert not df_cat.empty
    assert "Category" in df_cat.columns
    cats = df_cat["Category"].tolist()
    assert "GEN" in cats
    assert "OBC-NCL" in cats

    df_quota = aggregates["by_quota"]
    assert not df_quota.empty
    assert "Quota" in df_quota.columns

    df_year = aggregates["by_year"]
    assert not df_year.empty
    assert "Academic Year" in df_year.columns


def test_export_policy_report_xlsx():
    repo_root = Path(__file__).resolve().parents[2]
    students_csv = repo_root / "data" / "synthetic" / "students.csv"
    students_df = pd.read_csv(students_csv)

    xlsx_bytes = export_policy_report_xlsx(students_df, suppress=True, threshold=5)

    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 1000

    # Inspect the generated Excel workbook using openpyxl
    import io

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames

    assert "By State" in sheet_names
    assert "By Program" in sheet_names
    assert "By Category" in sheet_names
    assert "By Quota" in sheet_names
    assert "By Academic Year" in sheet_names

    # Check cell contents in 'By State' sheet
    ws_state = wb["By State"]
    headers = [cell.value for cell in ws_state[1]]
    assert "State" in headers
    assert "Total Students" in headers

    # Verify that cells were written
    assert ws_state.max_row >= 4
